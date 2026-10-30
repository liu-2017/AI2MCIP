#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""以 论文_优化稿v2.docx 为基础，按《中国环境监测》稿件格式要求生成投稿版 论文_投稿V1.docx。

主要处理：
  1) 页面与章节：删去"引言"标题（按期刊要求不设标题、不编号），章节号前移；
  2) 前置部分：题名/作者/单位/摘要/关键词/中图分类号 + 英文对应部分 + 脚注信息；
  3) 图表：按出现顺序重排图号，图表题按 小5号黑体 + 英文题，表改为三线表；
  4) 正文：5 号宋体、1.0 倍行距；各级标题字号按示例；
  5) 参考文献：按 GB/T 7714 顺序编码制重排，中文文献附英文题名；
  6) 去掉修订用红色字体与"附录 A 核对事项"。
"""
import copy
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from docx.text.paragraph import Paragraph

PAPER = '/home/ubuntu01/sda/AI2MCIP/paper'
SRC = os.path.join(PAPER, '论文_优化稿v2.docx')
DST = os.path.join(PAPER, '论文_投稿V1.docx')

BLACK = RGBColor(0, 0, 0)
SONG, HEI, KAI, TNR = '宋体', '黑体', '楷体', 'Times New Roman'

# 图号按"出现顺序"重排：旧号 -> 新号
FIG_MAP = {2: 1, 1: 2, 6: 3, 7: 4, 10: 5, 9: 6, 3: 7, 5: 8, 4: 9, 8: 10}

FIG_EN = {
    1: 'Model domain.',
    2: 'Schematic diagram of the five modelling chains and the AI-to-MCIP adapter.',
    3: 'Meteorological diagnostics: (a) mean diurnal cycle of PBLH; (b) daily PBLH; (c) daily 10 m wind speed; '
       '(d) daily ventilation index.',
    4: 'Mean vertical profiles of (a) potential temperature, (b) relative humidity and (c) PM2.5 over 14 cities '
       'on 30 October 2025.',
    5: 'Wall-clock time per simulated day.',
    6: 'Sensitivity experiments: (a) PBLH diurnal cycle under A1 and A2; (b) PM2.5 series; (c) KZMIN versus KZ0UT; '
       '(d) NMB and RMSE.',
    7: 'Hourly PM2.5 series: (a) 14-city mean; (b) Beijing, Shijiazhuang and Tianjin.',
    8: 'Statistical metrics: (a) scatter plot; (b) R and IOA; (c) NMB by phase; (d) RMSE by phase.',
    9: 'Spatial distribution of PM2.5 at 12:00 UTC on 30 October 2025.',
    10: 'Bias attribution: (a) NMB versus ventilation-index bias; (b) NMB versus PBLH bias; (c) chain-substitution '
        'effects; (d) daily mean bias.',
}

FIG_KEYS = ['模拟区域', '模拟链条与 AI-to-MCIP', '气象诊断量', '14 城市平均垂直廓线',
            '单位模拟日的墙钟时间对比', '诊断敏感性实验。', 'PM2.5逐时演变', '统计指标。',
            '空间分布', '偏差归因。']
TAB_KEYS = ['适配器的变量与来源', '计算环境与并行配置', '各环节墙钟时间分解',
            '诊断敏感性实验的总体指标', '5 条链条的 PM2.5模拟统计', '分阶段 NMB']

TAB_EN = {
    1: 'Variables provided by the adapter and their sources',
    2: 'Computing environment and parallel configuration',
    3: 'Wall-clock time decomposition (94 × 94 × 14 grid, min·d-1)',
    4: 'Overall metrics of the sensitivity experiments (14 cities × 216 h)',
    5: 'PM2.5 statistics of the five chains (14 cities × 216 h)',
    6: 'NMB (%) by episode phase',
}

REFS = [
    ('BI K, XIE L, ZHANG H, et al. Accurate medium-range global weather forecasting with 3D neural networks[J]. '
     'Nature, 2023, 619(7970): 533-538.', None),
    ('CHEN K, HAN T, GONG J, et al. FengWu: pushing the skillful global medium-range weather forecast beyond '
     '10 days lead[J/OL]. arXiv:2304.02948, 2023.', None),
    ('CHEN L, ZHONG X, ZHANG F, et al. FuXi: a cascade machine learning forecasting system for 15-day global '
     'weather forecast[J]. npj Climate and Atmospheric Science, 2023, 6: 190.', None),
    ('朱恩达，王亚强，赵妍，等. 东亚区域人工智能气象大模型预报技巧评估[J]. 应用气象学报，2024，35(6)：641-653.',
     'ZHU E D, WANG Y Q, ZHAO Y, et al. Evaluation of weather forecasts from AI big models over East Asia[J]. '
     'Journal of Applied Meteorological Science, 2024, 35(6): 641-653.'),
    ('SKAMAROCK W C, KLEMP J B, DUDHIA J, et al. A description of the Advanced Research WRF version 3[R]. '
     'Boulder: National Center for Atmospheric Research, 2008.', None),
    ('OTTE T L, PLEIM J E. The Meteorology-Chemistry Interface Processor (MCIP) for the CMAQ modeling system: '
     'updates through MCIPv3.4.1[J]. Geoscientific Model Development, 2010, 3(1): 243-256.', None),
    ('BYUN D, SCHERE K L. Review of the governing equations, computational algorithms, and other components of '
     'the Models-3 Community Multiscale Air Quality (CMAQ) modeling system[J]. Applied Mechanics Reviews, 2006, '
     '59(2): 51-77.', None),
    ('APPEL K W, NAPELENOK S L, FOLEY K M, et al. Description and evaluation of the Community Multiscale Air '
     'Quality (CMAQ) modeling system version 5.1[J]. Geoscientific Model Development, 2017, 10(4): 1703-1732.', None),
    ('KIM S M, KWON E, CHOI J H, et al. Application of machine learning-based wind corrections to coupled '
     'WRF-CMAQ simulations: a comparison of dynamic assimilation and static replacement[J]. Atmospheric '
     'Environment, 2026, 384: 122324.', None),
    ('张国斌，曹靖原，邱雄辉，等. 基于过程分析的京津冀地区气象条件变化对PM2.5浓度改善的影响[J]. 环境科学，'
     '2024，45(11)：6219-6228.',
     'ZHANG G B, CAO J Y, QIU X H, et al. Impact of change in meteorological conditions on PM2.5 air quality '
     'improvement in Beijing-Tianjin-Hebei region using process analysis[J]. Environmental Science, 2024, '
     '45(11): 6219-6228.'),
    ('CAMPBELL P C, TANG Y, LEE P, et al. Development and evaluation of an advanced National Air Quality '
     'Forecasting Capability using the NOAA Global Forecast System version 16[J]. Geoscientific Model '
     'Development, 2022, 15(8): 3281-3313.', None),
    ('TANG Y, CAMPBELL P C, LEE P, et al. Evaluation of the NAQFC driven by the NOAA Global Forecast System '
     '(version 16): comparison with the WRF-CMAQ during the summer 2019 FIREX-AQ campaign[J]. Geoscientific '
     'Model Development, 2022, 15(21): 7977-8004.', None),
    ('PLEIM J E. A combined local and nonlocal closure model for the atmospheric boundary layer. Part I: model '
     'description and testing[J]. Journal of Applied Meteorology and Climatology, 2007, 46(9): 1383-1395.', None),
    ('LI M, LIU H, GENG G, et al. Anthropogenic emission inventories in China: a review[J]. National Science '
     'Review, 2017, 4(6): 834-866.', None),
    ('宋小涵，燕丽，刘伟，等. 2015—2021年京津冀及周边地区PM2.5和臭氧复合污染时空特征分析[J]. 环境科学，2023，'
     '44(4)：1841-1851.',
     'SONG X H, YAN L, LIU W, et al. Spatiotemporal distribution characteristics of co-pollution of PM2.5 and '
     'ozone over BTH with surrounding area from 2015 to 2021[J]. Environmental Science, 2023, 44(4): 1841-1851.'),
    ('HERSBACH H, BELL B, BERRISFORD P, et al. The ERA5 global reanalysis[J]. Quarterly Journal of the Royal '
     'Meteorological Society, 2020, 146(730): 1999-2049.', None),
    ('JANJIĆ Z I. The step-mountain eta coordinate model: further developments of the convection, viscous '
     'sublayer, and turbulence closure schemes[J]. Monthly Weather Review, 1994, 122(5): 927-945.', None),
    ('李博文，张凯，高健，等. 京津冀及周边地区夏季和初秋气象条件对PM2.5和O3双高影响[J]. 环境科学，2024，'
     '45(8)：4375-4384.',
     'LI B W, ZHANG K, GAO J, et al. Effects of meteorological conditions in summer and early autumn on PM2.5 '
     'and O3 pollution in Beijing-Tianjin-Hebei and surrounding areas[J]. Environmental Science, 2024, '
     '45(8): 4375-4384.'),
    ('GU H, XU Y, PAN H, et al. A new framework for medium- to long-term PM2.5 predictions using AI-based '
     'high-resolution meteorological forecasts[J]. Chemical & Bio Engineering, 2026, 3(2): 219-230.', None),
    ('LIANG Z M, LI Y, WANG Y Q, et al. FengYuan: an end-to-end global weather forecasting model driven by '
     'multi-source observational data[J]. Journal of Meteorological Research, 2026, 40(4): 1120-1136.', None),
]


# ---------------------------------------------------------------- 基础工具
def delete_par(p):
    p._element.getparent().remove(p._element)


def set_run(run, cn=SONG, en=TNR, size=10.5, bold=False):
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK


def fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2, before=0, after=0,
            line=1.0, size=10.5):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Pt(size * indent_chars) if indent_chars else Pt(0)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def rewrite(p, text, cn=SONG, en=TNR, size=10.5, bold=False, **kw):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    set_run(r, cn, en, size, bold)
    fmt_par(p, size=size, **kw)


def insert_after(anchor, text, cn=SONG, en=TNR, size=10.5, bold=False, **kw):
    el = copy.deepcopy(anchor._p)
    anchor._p.addnext(el)
    p = Paragraph(el, anchor._parent)
    rewrite(p, text, cn, en, size, bold, **kw)
    return p


def find(doc, sub, nth=0):
    hits = [p for p in doc.paragraphs if sub in p.text]
    if len(hits) <= nth:
        raise KeyError('未找到：%r' % sub)
    return hits[nth]


def fig_renumber(text):
    def rep(m):
        n = int(m.group(1))
        return '图 %d%s' % (FIG_MAP.get(n, n), m.group(2) or '')
    return re.sub(r'图\s*(\d+)([a-d]?)', rep, text)


def sec_renumber(text):
    text = re.sub(r'(\d)\.(\d)\s*节', lambda m: '%d.%d 节' % (int(m.group(1)) - 1, int(m.group(2)))
                  if int(m.group(1)) >= 2 else m.group(0), text)
    return text


# ---------------------------------------------------------------- 主流程
def main():
    doc = Document(SRC)
    log = []

    # 1) 删除：红色说明、附录 A
    for p in list(doc.paragraphs):
        if p.text.strip().startswith('说明：本稿中红色字体'):
            delete_par(p)
            log.append('删除修订说明行')
    # 参考文献标题及其后所有内容（含旧条目与附录 A）整体删除，稍后重建
    paras = doc.paragraphs
    idx = next((i for i, p in enumerate(paras) if p.text.strip() == '参考文献'), None)
    if idx is not None:
        for p in paras[idx:]:
            delete_par(p)
        log.append('删除旧参考文献与附录 A（%d 段）' % (len(paras) - idx))

    # 2) 图号按出现顺序重排（全文，含图题）
    for p in doc.paragraphs:
        t = p.text
        if '图' in t:
            nt = fig_renumber(t)
            if nt != t:
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                p.add_run(nt)

    # 3) 章节号前移（删去"引言"后，2→1、3→2、4→3、5→4）
    for p in list(doc.paragraphs):
        t = p.text.strip()
        m = re.match(r'^(\d)(\s+)(.+)$', t)
        if m and p.style.name.startswith('Heading'):
            old = int(m.group(1))
            new = old - 1 if old >= 2 else old
            txt = '%d%s%s' % (new, m.group(2), m.group(3))
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(txt)
        m2 = re.match(r'^(\d)\.(\d)(\s+)(.+)$', t)
        if m2 and p.style.name.startswith('Heading'):
            old = int(m2.group(1))
            new = old - 1 if old >= 2 else old
            txt = '%d.%s%s%s' % (new, m2.group(2), m2.group(3), m2.group(4))
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(txt)
    # 交叉引用
    for p in doc.paragraphs:
        t = p.text
        nt = sec_renumber(t)
        if nt != t:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(nt)

    # 4) 删除"1 引言""摘要""Abstract"标题与旧"参考文献"标题（稍后重建）
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t in ('1 引言', '1  引言', '摘要', 'Abstract'):
            # 删除前先记录，参考文献内容稍后重建时整体删
            delete_par(p)
    log.append('删除"引言/摘要/Abstract/参考文献"标题')

    # 6) 前置部分
    title = doc.paragraphs[0]
    rewrite(title, title.text.strip(), cn=HEI, size=18, align=WD_ALIGN_PARAGRAPH.CENTER,
            indent_chars=0, before=0, after=6, line=1.2)
    log.append('中文题名：小2号黑体')

    # 6.1 作者与单位（待补）
    p_author = insert_after(title, '作者姓名（待补）', cn=KAI, size=12,
                            align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=6, after=2, line=1.0)
    p_aff = insert_after(p_author, '1. 单位名称（待补），城市 邮编（待补）', cn=SONG, size=9,
                         align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=0, after=6, line=1.0)

    # 6.2 中文摘要：把原摘要段落改为"摘要：…"
    abs_cn = None
    for p in doc.paragraphs:
        if p.text.startswith('全球AI气象模型为区域空气质量模拟'):
            abs_cn = p
            break
    if abs_cn is None:
        raise KeyError('未找到中文摘要段')
    body = abs_cn.text.strip()
    for r in list(abs_cn.runs):
        r._element.getparent().remove(r._element)
    r1 = abs_cn.add_run('摘  要：')
    set_run(r1, HEI, TNR, 9, False)
    r2 = abs_cn.add_run(body)
    set_run(r2, KAI, TNR, 9, False)
    fmt_par(abs_cn, indent_chars=0, before=6, after=2, line=1.1, size=9)

    # 6.3 关键词（标签黑体、内容楷体）
    kw_cn = find(doc, '关键词：')
    content = kw_cn.text.replace('关键词：', '').strip()
    for r in list(kw_cn.runs):
        r._element.getparent().remove(r._element)
    r1 = kw_cn.add_run('关键词：')
    set_run(r1, HEI, TNR, 9, False)
    r2 = kw_cn.add_run(content)
    set_run(r2, KAI, TNR, 9, False)
    fmt_par(kw_cn, indent_chars=0, before=0, after=2, line=1.1, size=9)

    # 6.4 中图分类号
    clc = insert_after(kw_cn, '中图分类号：X51', cn=KAI, size=9,
                       align=WD_ALIGN_PARAGRAPH.LEFT, indent_chars=0, before=0, after=8, line=1.1)
    r = clc.runs[0]
    set_run(r, HEI, TNR, 9, False)
    r2 = clc.add_run('X51')
    set_run(r2, KAI, TNR, 9, False)
    r.text = '中图分类号：'
    log.append('新增中图分类号')

    # 6.5 英文题名/作者/单位/Abstract/Keywords
    en_title = find(doc, 'A Standardized AI-to-MCIP Adapter')
    rewrite(en_title, en_title.text.strip(), cn=HEI, en=TNR, size=9, bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=8, after=2, line=1.1)
    p_en_auth = insert_after(en_title, 'Author Name (TBD)', cn=SONG, en=TNR, size=9,
                             align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=0, after=0, line=1.1)
    p_en_aff = insert_after(p_en_auth, '1. Affiliation (TBD), City Postcode, China', cn=SONG, en=TNR, size=9,
                            align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=0, after=4, line=1.1)
    en_abs = find(doc, 'Global AI weather models offer an alternative')
    body_en = en_abs.text.strip()
    for r in list(en_abs.runs):
        r._element.getparent().remove(r._element)
    r1 = en_abs.add_run('Abstract: ')
    set_run(r1, TNR, TNR, 9, True)
    r2 = en_abs.add_run(body_en)
    set_run(r2, TNR, TNR, 9, False)
    fmt_par(en_abs, indent_chars=0, before=2, after=2, line=1.1, size=9)

    en_kw = find(doc, 'Keywords:')
    content_en = en_kw.text.replace('Keywords:', '').strip()
    for r in list(en_kw.runs):
        r._element.getparent().remove(r._element)
    r1 = en_kw.add_run('Keywords: ')
    set_run(r1, TNR, TNR, 9, True)
    r2 = en_kw.add_run(content_en)
    set_run(r2, TNR, TNR, 9, False)
    fmt_par(en_kw, indent_chars=0, before=0, after=8, line=1.1, size=9)

    # 6.6 脚注信息（收稿日期等）
    feet = ['收稿日期：××××-××-××；修订日期：××××-××-××',
            '基金项目：××（项目编号：××）',
            '第一作者简介：×××（19××—），男/女，硕士/博士，×级工程师，E-mail：××@××.××',
            '通信作者：×××，E-mail：××@××.××']
    prev = en_kw
    for i, line in enumerate(feet):
        prev = insert_after(prev, line, cn=SONG, size=9, align=WD_ALIGN_PARAGRAPH.LEFT,
                            indent_chars=0, before=(4 if i == 0 else 0),
                            after=(10 if i == len(feet) - 1 else 0), line=1.05)
    log.append('新增脚注信息（4 行）')

    # 6.7 过长标题精简（期刊要求章节标题一般 ≤15 字）
    SHORT = {
        '2.1 接口输出被 MCIP/CMAQ 解析与执行': '2.1 接口执行情况',
        '2.4 诊断模块的敏感性检验：诊断选择是否改变了主要结论？': '2.4 诊断模块的敏感性检验',
        '3 PM2.5重污染过程模拟结果与偏差归因': '3 模拟结果与偏差归因',
        '3.3 与已有 AI 驱动空气质量工作的对话': '3.3 与已有工作的对比',
    }
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in SHORT:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(SHORT[t])
    log.append('精简超长章节标题 %d 处' % len(SHORT))

    # 7) 正文格式 + 图表题
    fig_cap = 0
    tab_cap = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        st = p.style.name
        if not t:
            continue
        if st.startswith('Heading 1'):
            rewrite(p, t, cn=SONG, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
                    indent_chars=0, before=12, after=6, line=1.2)
        elif st.startswith('Heading 2'):
            rewrite(p, t, cn=HEI, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
                    indent_chars=0, before=8, after=4, line=1.2)
        elif st.startswith('Heading 3'):
            rewrite(p, t, cn=SONG, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                    indent_chars=0, before=6, after=3, line=1.2)
        elif re.match(r'^图\s*\d+\s', t) and any(k in t for k in FIG_KEYS):
            rewrite(p, t, cn=HEI, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
                    indent_chars=0, before=4, after=2, line=1.1)
            m = re.match(r'^图\s*(\d+)\s*(.*)$', t)
            n = int(m.group(1)) if m else None
            if n in FIG_EN:
                insert_after(p, 'Fig. %d  %s' % (n, FIG_EN[n]), cn=SONG, en=TNR, size=9,
                             align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=0, after=6, line=1.1)
            fig_cap += 1
        elif re.match(r'^表\s*\d+\s', t) and any(k in t for k in TAB_KEYS):
            rewrite(p, t, cn=HEI, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
                    indent_chars=0, before=8, after=2, line=1.1)
            m = re.match(r'^表\s*(\d+)\s*(.*)$', t)
            n = int(m.group(1)) if m else None
            if n in TAB_EN:
                insert_after(p, 'Table %d  %s' % (n, TAB_EN[n]), cn=SONG, en=TNR, size=9,
                             align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, before=0, after=4, line=1.1)
            tab_cap += 1
        elif t.startswith('注：'):
            rewrite(p, t, cn=SONG, size=7.5, align=WD_ALIGN_PARAGRAPH.LEFT,
                    indent_chars=0, before=2, after=8, line=1.0)
        else:
            rewrite(p, t, cn=SONG, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                    indent_chars=2, before=0, after=0, line=1.0)
    log.append('正文格式：5号宋体/1.0 倍行距；图题 %d、表题 %d' % (fig_cap, tab_cap))

    # 8) 表格：三线表 + 9 磅宋体
    for tb in doc.tables:
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        _three_line(tb)
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        set_run(r, SONG, TNR, 9, False)
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
    log.append('表格：三线表 + 小5号')

    # 9) 参考文献
    anchor = find(doc, '数据与代码可用性')
    # 在"数据与代码可用性"整段之后追加参考文献
    last = doc.paragraphs[-1]
    h = insert_after(last, '参考文献（References）：', cn=SONG, en=TNR, size=10.5, bold=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT, indent_chars=0, before=12, after=6, line=1.2)
    prev = h
    for i, (cn_txt, en_txt) in enumerate(REFS, 1):
        prev = insert_after(prev, '[%d]  %s' % (i, cn_txt), cn=SONG, en=TNR, size=9,
                            align=WD_ALIGN_PARAGRAPH.LEFT, indent_chars=0, before=0, after=1, line=1.0)
        prev.paragraph_format.left_indent = Cm(0.75)
        prev.paragraph_format.first_line_indent = Cm(-0.75)
        if en_txt:
            prev = insert_after(prev, en_txt, cn=SONG, en=TNR, size=9,
                                align=WD_ALIGN_PARAGRAPH.LEFT, indent_chars=0, before=0, after=1, line=1.0)
            prev.paragraph_format.left_indent = Cm(0.75)
    log.append('参考文献：%d 条（GB/T 7714，中文附英文）' % len(REFS))

    # 10) 全部文字改黑（去掉修订红）
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = BLACK
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = BLACK

    # 11) 页面设置
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for attr in ('left_margin', 'right_margin'):
        setattr(s, attr, Cm(2.5))
    s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.5)
    log.append('页面：A4，四周 2.5 cm')

    doc.save(DST)
    print('saved:', DST)
    for i, x in enumerate(log, 1):
        print('  %2d. %s' % (i, x))


def _three_line(tb):
    """三线表：上下粗线、表头下细线，无竖线。"""
    tbl = tb._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for tag, sz in (('top', 12), ('left', 0), ('bottom', 12), ('right', 0),
                    ('insideH', 0), ('insideV', 0)):
        el = OxmlElement('w:%s' % tag)
        el.set(qn('w:val'), 'single' if sz else 'none')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)
    # 表头下横线
    if len(tb.rows) > 1:
        for c in tb.rows[0].cells:
            tcPr = c._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), '000000')
            b.append(bottom)
            tcPr.append(b)


if __name__ == '__main__':
    main()
